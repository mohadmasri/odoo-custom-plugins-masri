import base64
import io
import re
from datetime import date, datetime, timedelta

from markupsafe import Markup, escape

from odoo import api, fields, models
from odoo.exceptions import AccessError, UserError, ValidationError


def normalize_account_key(value):
    """توحيد شكل اسم/رمز الحساب لمقارنته: إزالة الفراغات الزائدة وتوحيد الحالة."""
    return ' '.join((value or '').split()).casefold()


def normalize_digits(value):
    """تحويل الأرقام الهندية (٠-٩) والفارسية (۰-۹) إلى أرقام لاتينية."""
    return (value or '').translate(str.maketrans('٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹', '01234567890123456789'))


# أسماء الأشهر المقبولة في عمود "شهر دفتر الأستاذ" عند الاستيراد
ARABIC_MONTHS = {
    name: number
    for number, names in enumerate([
        ('يناير', 'كانون الثاني', 'jan', 'january'),
        ('فبراير', 'شباط', 'feb', 'february'),
        ('مارس', 'آذار', 'اذار', 'mar', 'march'),
        ('أبريل', 'ابريل', 'نيسان', 'apr', 'april'),
        ('مايو', 'أيار', 'ايار', 'may'),
        ('يونيو', 'حزيران', 'jun', 'june'),
        ('يوليو', 'تموز', 'jul', 'july'),
        ('أغسطس', 'اغسطس', 'آب', 'اب', 'aug', 'august'),
        ('سبتمبر', 'أيلول', 'ايلول', 'sep', 'september'),
        ('أكتوبر', 'اكتوبر', 'تشرين الأول', 'تشرين الاول', 'oct', 'october'),
        ('نوفمبر', 'تشرين الثاني', 'nov', 'november'),
        ('ديسمبر', 'كانون الأول', 'كانون الاول', 'dec', 'december'),
    ], start=1)
    for name in names
}

MOVE_TYPES = [
    ('entry', 'قيد محاسبي'),
    ('receipt', 'سند قبض'),
]

# تفقيط المبالغ بالعربية
ARABIC_ONES = ['', 'واحد', 'اثنان', 'ثلاثة', 'أربعة', 'خمسة', 'ستة', 'سبعة', 'ثمانية', 'تسعة',
               'عشرة', 'أحد عشر', 'اثنا عشر', 'ثلاثة عشر', 'أربعة عشر', 'خمسة عشر',
               'ستة عشر', 'سبعة عشر', 'ثمانية عشر', 'تسعة عشر']
ARABIC_TENS = ['', '', 'عشرون', 'ثلاثون', 'أربعون', 'خمسون', 'ستون', 'سبعون', 'ثمانون', 'تسعون']
ARABIC_HUNDREDS = ['', 'مائة', 'مائتان', 'ثلاثمائة', 'أربعمائة', 'خمسمائة', 'ستمائة',
                   'سبعمائة', 'ثمانمائة', 'تسعمائة']
ARABIC_SCALES = [('', '', ''), ('ألف', 'ألفان', 'آلاف'), ('مليون', 'مليونان', 'ملايين'),
                 ('مليار', 'ملياران', 'مليارات')]
# أسماء العملات ووحداتها الصغرى بالعربية (وإلا استُخدم اسم العملة كما هو)
ARABIC_CURRENCIES = {
    # الرمز: (مفرد، جمع، الوحدة الصغرى مفرد، الوحدة الصغرى جمع)
    'JOD': ('دينار أردني', 'دنانير أردنية', 'فلس', 'فلوس'),
    'USD': ('دولار أمريكي', 'دولارات أمريكية', 'سنت', 'سنتات'),
    'EUR': ('يورو', 'يوروهات', 'سنت', 'سنتات'),
    'SAR': ('ريال سعودي', 'ريالات سعودية', 'هللة', 'هللات'),
    'AED': ('درهم إماراتي', 'دراهم إماراتية', 'فلس', 'فلوس'),
    'EGP': ('جنيه مصري', 'جنيهات مصرية', 'قرش', 'قروش'),
    'ILS': ('شيكل', 'شواكل', 'أغورة', 'أغورات'),
}


def arabic_currency_label(count, singular, plural):
    """الجمع العربي: من 3 إلى 10 يُجمع، وما عداه يبقى مفرداً."""
    return plural if 3 <= count % 100 <= 10 else singular


def arabic_number_to_words(number):
    """يحوّل عدداً صحيحاً إلى كلمات عربية."""
    number = int(number)
    if number == 0:
        return 'صفر'

    def three(value):
        parts = []
        if value >= 100:
            parts.append(ARABIC_HUNDREDS[value // 100])
            value %= 100
        if value >= 20:
            unit = value % 10
            if unit:
                parts.append(ARABIC_ONES[unit])
            parts.append(ARABIC_TENS[value // 10])
        elif value:
            parts.append(ARABIC_ONES[value])
        return ' و'.join(parts)

    groups = []
    scale = 0
    while number and scale < len(ARABIC_SCALES):
        groups.append((scale, number % 1000))
        number //= 1000
        scale += 1

    words = []
    for scale, value in reversed(groups):
        if not value:
            continue
        if scale == 0:
            words.append(three(value))
            continue
        singular, dual, plural = ARABIC_SCALES[scale]
        if value == 1:
            words.append(singular)
        elif value == 2:
            words.append(dual)
        elif 3 <= value <= 10:
            words.append(f'{three(value)} {plural}')
        else:
            words.append(f'{three(value)} {singular}')
    return ' و'.join(words)


LEDGER_MONTH_SELECTION = [
    ('1', '1 - يناير'), ('2', '2 - فبراير'), ('3', '3 - مارس'), ('4', '4 - أبريل'),
    ('5', '5 - مايو'), ('6', '6 - يونيو'), ('7', '7 - يوليو'), ('8', '8 - أغسطس'),
    ('9', '9 - سبتمبر'), ('10', '10 - أكتوبر'), ('11', '11 - نوفمبر'), ('12', '12 - ديسمبر'),
]


class MyAccountingMove(models.Model):
    _name = 'myaccounting.move'
    _description = 'قيد محاسبي'
    _order = 'sort_key desc, id desc'
    _inherit = ['mail.thread', 'mail.activity.mixin']

    name = fields.Char(string='رقم القيد', required=True, copy=False, tracking=True,
                        default=lambda self: self._get_default_name())
    date = fields.Date(string='التاريخ', required=True, tracking=True,
                       default=lambda self: self._get_default_date())
    ref = fields.Char(string='المرجع')
    journal = fields.Char(
        string='اليومية', default='القيود اليدوية',
        help='يمكن اختيار أكثر من يومية للقيد الواحد (مثل: ايرادات، رواتب).')

    ledger_month = fields.Selection(
        LEDGER_MONTH_SELECTION, string='شهر دفتر الأستاذ', required=True, tracking=True,
        default=lambda self: self._get_default_ledger_period()[1],
        help='الشهر الذي سيظهر فيه هذا القيد ضمن دفتر الأستاذ العام، بغض النظر عن تاريخ القيد أعلاه.',
    )
    ledger_year = fields.Integer(
        string='سنة دفتر الأستاذ', required=True, tracking=True,
        default=lambda self: self._get_default_ledger_period()[0],
    )
    ledger_period_label = fields.Char(string='الشهر المحاسبي', compute='_compute_ledger_period_label', store=True)
    # مفتاح الترتيب الموحّد لكل الشاشات: السنة ثم الشهر ثم النوع ثم رقم القيد
    # (11/8 قبل 12/8 مهما كان ترتيب الإدخال، والقيود القديمة المُدخلة لاحقاً تأخذ مكانها)
    sort_key = fields.Char(string='مفتاح الترتيب', compute='_compute_sort_key', store=True, index=True)

    move_type = fields.Selection(
        MOVE_TYPES, string='النوع', default='entry', required=True, copy=True,
        help='سند القبض هو قيد محاسبي بترقيم مستقل وشكل طباعة مختلف.')

    state = fields.Selection([
        ('draft', 'مسودة'),
        ('incomplete', 'غير مكتمل'),
        ('posted', 'مرحّل'),
    ], string='الحالة', default='draft', tracking=True, copy=False,
        help='"غير مكتمل" يعني أن القيد يحتوي بنوداً لم يُعرف حسابها عند الاستيراد من ملف Excel، '
             'ويجب اختيار الحساب يدوياً قبل الترحيل.')

    line_ids = fields.One2many('myaccounting.move.line', 'move_id', string='بنود القيد', copy=True)

    total_debit = fields.Float(string='إجمالي المدين', compute='_compute_totals', store=True, digits=(16, 3))
    total_credit = fields.Float(string='إجمالي الدائن', compute='_compute_totals', store=True, digits=(16, 3))
    is_balanced = fields.Boolean(string='متوازن', compute='_compute_totals', store=True)
    currency_id = fields.Many2one('res.currency', string='العملة',
                                   default=lambda self: self.env.company.currency_id)
    company_id = fields.Many2one('res.company', string='الشركة', default=lambda self: self.env.company)

    # المستندات المرفقة بالقيد (صورة الفاتورة أو الشيك أو أي مستند مؤيد)
    attachment_ids = fields.Many2many(
        'ir.attachment', 'myaccounting_move_attachment_rel', 'move_id', 'attachment_id',
        string='المرفقات')
    attachment_count = fields.Integer(string='عدد المرفقات', compute='_compute_attachment_count')

    @api.depends('attachment_ids')
    def _compute_attachment_count(self):
        for move in self:
            move.attachment_count = len(move.attachment_ids)

    # ملاحظات الاستيراد: الأخطاء والتصحيحات التلقائية التي حدثت عند استيراد القيد
    # من Excel. تُنشر أيضاً كـ"ملاحظة" في المحادثة، وتبقى ظاهرة ومميّزة حتى تُراجَع.
    import_notes = fields.Html(string='ملاحظات الاستيراد', readonly=True, copy=False, sanitize=True)
    import_notes_reviewed = fields.Boolean(string='تمت مراجعة ملاحظات الاستيراد', copy=False)
    has_import_notes = fields.Boolean(string='لديه ملاحظات استيراد', compute='_compute_has_import_notes',
                                      store=True)

    @api.depends('import_notes', 'import_notes_reviewed')
    def _compute_has_import_notes(self):
        for move in self:
            move.has_import_notes = bool(move.import_notes) and not move.import_notes_reviewed

    def action_mark_import_notes_reviewed(self):
        self.write({'import_notes_reviewed': True})
        for move in self:
            move.message_post(body='تمت مراجعة ملاحظات الاستيراد.', subtype_xmlid='mail.mt_note')
        return True

    _name_company_uniq = models.Constraint(
        'unique(name, company_id)',
        'رقم القيد هذا مستخدم بالفعل، الرجاء اختيار رقم آخر.',
    )

    @api.model
    def _increment_name(self, name):
        """يزيد أول رقم يظهر في النص بمقدار 1 مع الحفاظ على باقي التنسيق.
        مثال: '1/10' تصبح '2/10'."""
        match = re.search(r'\d+', name)
        if not match:
            return f'{name}-1'
        num = str(int(match.group()) + 1).zfill(len(match.group()))
        return name[:match.start()] + num + name[match.end():]

    @api.model
    def _get_default_name(self):
        # لكل نوع ترقيمه المستقل: يؤخذ رقم آخر سجل من نفس النوع ويُزاد واحداً
        move_type = self.env.context.get('default_move_type') or 'entry'
        last_move = self.search([('move_type', '=', move_type)], order='id desc', limit=1)
        candidate = self._increment_name(last_move.name) if last_move and last_move.name else '1'
        guard = 0
        while guard < 1000 and self.search_count([('name', '=', candidate)]):
            candidate = self._increment_name(candidate)
            guard += 1
        return candidate

    # ========================================================================
    # قوالب القيود: قائمة اليوميات تحت "قيد جديد"
    # ========================================================================

    def _register_hook(self):
        """يبني/يحدّث قوائم قوالب اليوميات عند كل إقلاع للخادم (وبعد كل تحديث)."""
        super()._register_hook()
        try:
            self.env['myaccounting.move']._sync_journal_template_menus()
            self.env.cr.commit()
        except Exception:  # noqa: BLE001 - لا نمنع إقلاع الخادم بسبب القوائم
            self.env.cr.rollback()

    # ========================================================================
    # قائمة التذكير في الصفحة الرئيسية (مرتبطة بتطبيق المهام To-do في أودو)
    # ========================================================================

    TODO_DOMAIN_OPEN = [('project_id', '=', False), ('state', 'not in', ['1_done', '1_canceled'])]

    @api.model
    def get_home_todos(self):
        """مهام المستخدم الشخصية غير المنجزة (نفس مهام تطبيق To-do)."""
        try:
            tasks = self.env['project.task'].search(
                self.TODO_DOMAIN_OPEN + [('user_ids', 'in', [self.env.uid])],
                order='priority desc, id desc', limit=50)
        except AccessError:
            return {'allowed': False, 'items': []}
        return {
            'allowed': True,
            'items': [{'id': task.id, 'name': task.name, 'starred': task.priority == '1'} for task in tasks],
        }

    @api.model
    def add_home_todo(self, name):
        name = (name or '').strip()
        if not name:
            raise UserError('اكتب نص التذكير أولاً.')
        self.env['project.task'].create({'name': name, 'user_ids': [(4, self.env.uid)]})
        return self.get_home_todos()

    @api.model
    def done_home_todo(self, task_id):
        task = self.env['project.task'].browse(task_id).exists()
        if task and self.env.uid in task.user_ids.ids:
            task.state = '1_done'
        return self.get_home_todos()

    @api.model
    @api.model
    def split_journals(self, text):
        """اليومية قد تحوي أكثر من اسم: "ايرادات، رواتب" → ['ايرادات', 'رواتب']."""
        return [part.strip() for part in re.split(r'[،,+]', text or '') if part.strip()]

    def has_journal(self, name):
        self.ensure_one()
        return name in self.split_journals(self.journal)

    def action_new_from_journal(self, journal):
        """ينشئ قيداً جديداً مطابقاً لبنود آخر قيد في نفس اليومية،
        برقم وتاريخ وشهر أستاذ يتبع آخر ما وصلت إليه القيود، ويفتحه للتعديل."""
        # آخر قيد تحوي يومياته هذه اليومية (قد يحمل القيد أكثر من يومية)
        template = next((move for move in self.search([('journal', 'ilike', journal)], order='id desc')
                         if move.has_journal(journal)), self.browse())
        if not template:
            raise UserError(f'لا يوجد قيد سابق في اليومية "{journal}" لاستخدامه كقالب.')
        ledger_year, ledger_month = self._get_default_ledger_period()
        new_move = template.copy({
            'date': self._get_default_date(),
            'ledger_year': ledger_year,
            'ledger_month': ledger_month,
        })
        new_move.message_post(
            body=f'أُنشئ هذا القيد من قالب اليومية "{journal}" نسخاً عن القيد {template.name}.',
            subtype_xmlid='mail.mt_note')
        return {
            'type': 'ir.actions.act_window',
            'res_model': 'myaccounting.move',
            'res_id': new_move.id,
            'views': [[False, 'form']],
            'target': 'current',
        }

    @api.model_create_multi
    def create(self, vals_list):
        moves = super().create(vals_list)
        moves._ensure_journal_menus()
        return moves

    def write(self, vals):
        res = super().write(vals)
        if 'journal' in vals:
            self._ensure_journal_menus()
        return res

    def unlink(self):
        journals = {name for move in self for name in self.split_journals(move.journal)}
        res = super().unlink()
        if journals:
            self._sync_journal_template_menus()
        return res

    def _ensure_journal_menus(self):
        """يسجّل أي يومية جديدة ويحدّث القائمة عند الحاجة."""
        journals = {name for move in self for name in self.split_journals(move.journal)}
        if not journals:
            return
        Journal = self.env['myaccounting.journal']
        if Journal.search_count([('name', 'in', list(journals))]) != len(journals):
            self._sync_journal_template_menus()

    @api.model
    def _journal_menu_xmlid(self, journal):
        digest = re.sub(r'[^a-z0-9]+', '_', journal.strip().casefold())
        return f'journal_tpl_{abs(hash(journal)) % (10 ** 8)}_{digest[:20]}'

    @api.model
    def _sync_journal_template_menus(self):
        """يحدّث قائمة القوالب تحت "قيد جديد" لتطابق اليوميات الموجودة فعلاً."""
        parent = self.env.ref('my_accounting.menu_myaccounting_move_new', raise_if_not_found=False)
        if not parent:
            return
        Menu = self.env['ir.ui.menu'].sudo()
        Server = self.env['ir.actions.server'].sudo()
        model = self.env['ir.model']._get('myaccounting.move')
        empty_menu = self.env.ref('my_accounting.menu_myaccounting_move_new_empty', raise_if_not_found=False)

        Journal = self.env['myaccounting.journal']
        Journal._sync_journals()
        journals = [
            journal.name for journal in Journal.search([('show_in_menu', '=', True)])
            if journal.move_count
        ]
        existing = Menu.search([('parent_id', '=', parent.id)])
        by_name = {menu.name: menu for menu in existing if not empty_menu or menu.id != empty_menu.id}

        for index, journal in enumerate(journals, start=10):
            menu = by_name.pop(journal, None)
            if menu:
                menu.sequence = index
                continue
            action = Server.create({
                'name': journal,
                'model_id': model.id,
                'state': 'code',
                'code': f'action = env["myaccounting.move"].action_new_from_journal({journal!r})',
            })
            Menu.create({
                'name': journal,
                'parent_id': parent.id,
                'sequence': index,
                'action': f'ir.actions.server,{action.id}',
                'group_ids': [(6, 0, parent.group_ids.ids)],
            })
        # يوميات لم تعد موجودة: تُحذف قائمتها
        for menu in by_name.values():
            menu.unlink()

    def receipt_amount_parts(self, amount=None):
        """يقسم المبلغ إلى (دينار، فلس) لعرضه في عمودين كما في نموذج السند."""
        self.ensure_one()
        currency = self.currency_id or self.env.company.currency_id
        decimals = currency.decimal_places or 2
        value = round(self.total_debit if amount is None else amount, decimals)
        units = int(value)
        return {'units': units, 'subunits': int(round((value - units) * (10 ** decimals)))}

    def receipt_payer(self):
        """"وصلني من السادة": الجهة الدافعة، وهي حساب الطرف الدائن في السند."""
        self.ensure_one()
        credit_lines = self.line_ids.filtered(lambda line: line.credit)
        if credit_lines:
            return credit_lines[0].account_label or credit_lines[0].account_id.name or ''
        return self.ref or ''

    def receipt_detail_lines(self):
        """بنود التفاصيل في السند: الأسطر المدينة (المبالغ المقبوضة)."""
        self.ensure_one()
        debit_lines = self.line_ids.filtered(lambda line: line.debit)
        return debit_lines or self.line_ids

    def currency_label_ar(self):
        """اسم العملة بالعربية حسب مبلغ السند (للطباعة)."""
        self.ensure_one()
        currency = self.currency_id or self.env.company.currency_id
        names = ARABIC_CURRENCIES.get(currency.name)
        if not names:
            return currency.name or ''
        total = int(round(self.total_debit or self.total_credit or 0.0))
        return arabic_currency_label(total, names[0], names[1])

    def amount_in_words(self):
        """المبلغ الإجمالي للسند مكتوباً بالكلمات بالعربية."""
        self.ensure_one()
        currency = self.currency_id or self.env.company.currency_id
        decimals = currency.decimal_places or 2
        names = ARABIC_CURRENCIES.get(currency.name)
        if not names:
            names = (currency.name or '', currency.name or '', '', '')
        unit_one, unit_many, sub_one, sub_many = names
        total = round(self.total_debit or self.total_credit or 0.0, decimals)
        units = int(total)
        subunits = int(round((total - units) * (10 ** decimals)))
        text = f'{arabic_number_to_words(units)} {arabic_currency_label(units, unit_one, unit_many)}'.strip()
        if subunits and sub_one:
            text += (f' و{arabic_number_to_words(subunits)} '
                     f'{arabic_currency_label(subunits, sub_one, sub_many)}')
        return f'{text} فقط لا غير'

    @api.model
    def _get_default_date(self):
        """تاريخ القيد الجديد = تاريخ آخر قيد مُدخَل (وتاريخ اليوم إن لم يوجد أي قيد).
        يبقى الحقل قابلاً للتعديل يدوياً كالمعتاد."""
        last_move = self.search([], order='id desc', limit=1)
        return last_move.date or fields.Date.context_today(self)

    @api.model
    def _get_default_ledger_period(self):
        last_move = self.search([], order='id desc', limit=1)
        if last_move and last_move.ledger_month and last_move.ledger_year:
            return last_move.ledger_year, last_move.ledger_month
        today = fields.Date.context_today(self)
        return today.year, str(today.month)

    @api.depends('name', 'ledger_year', 'ledger_month', 'move_type')
    def _compute_sort_key(self):
        Account = self.env['myaccounting.account']
        for move in self:
            year, month, kind, number, rest = Account._move_sort_key(move)
            move.sort_key = f'{year:04d}{month:02d}{kind}{number:010d}{rest}'

    @api.depends('ledger_month', 'ledger_year')
    def _compute_ledger_period_label(self):
        for move in self:
            if move.ledger_month and move.ledger_year:
                move.ledger_period_label = f"{int(move.ledger_month):02d}/{move.ledger_year}"
            else:
                move.ledger_period_label = ''

    @api.depends('line_ids.debit', 'line_ids.credit')
    def _compute_totals(self):
        for move in self:
            move.total_debit = sum(move.line_ids.mapped('debit'))
            move.total_credit = sum(move.line_ids.mapped('credit'))
            move.is_balanced = round(move.total_debit - move.total_credit, 3) == 0.0

    def _sync_state_from_lines(self):
        """يحدّث حالة القيد بين "مسودة" و"غير مكتمل" حسب وجود بنود بلا حساب.
        لا يمسّ القيود المرحّلة."""
        for move in self:
            if move.state == 'posted':
                continue
            new_state = 'incomplete' if any(not line.account_id for line in move.line_ids) else 'draft'
            if move.state != new_state:
                move.state = new_state

    def action_post(self):
        for move in self:
            if not move.line_ids:
                raise UserError('لا يمكن ترحيل قيد بدون بنود.')
            missing = move.line_ids.filtered(lambda line: not line.account_id)
            if missing:
                names = '، '.join(sorted({line.pending_account_name or '—' for line in missing}))
                raise UserError(
                    f'لا يمكن ترحيل القيد "{move.name}" لأنه يحتوي بنوداً بلا حساب محدّد '
                    f'(أسماء الحسابات في الملف: {names}).\n'
                    'اختر الحساب الصحيح لكل بند أولاً.'
                )
            if not move.is_balanced:
                raise UserError(
                    f'القيد غير متوازن: إجمالي المدين {move.total_debit} لا يساوي إجمالي الدائن {move.total_credit}.'
                )
            if not move.name:
                raise UserError('يجب إدخال رقم للقيد قبل الترحيل.')
            move.state = 'posted'

    def action_post_and_next(self):
        """نفس الترحيل؛ الانتقال للقيد التالي يتم في الواجهة (move_form.js)
        بعد نجاح الترحيل فقط."""
        return self.action_post()

    def action_reset_to_draft(self):
        self.write({'state': 'draft'})
        self._sync_state_from_lines()

    def unlink(self):
        if not self.env.context.get('force_delete'):
            for move in self:
                if move.state == 'posted':
                    raise UserError('لا يمكن حذف قيد مرحّل. أعده إلى مسودة أولاً.')
        return super().unlink()

    @api.model
    def get_general_ledger_matrix(self, year, month, states=None):
        """دفتر الأستاذ العام: صف لكل قيد، وعمودا مدين/دائن لكل حساب رئيسي (جذري) دائماً،
        مع تجميع حركات أي حساب فرعي تحت حسابه الرئيسي.
        يعتمد التصنيف على شهر/سنة دفتر الأستاذ (ledger_month/ledger_year) للقيد
        وليس على تاريخه الفعلي، لدعم القيود التي تُسجَّل في شهر مختلف عن تاريخها."""
        # الحالات المطلوبة: مرحّل فقط افتراضياً (كما كان)، أو ما يختاره المستخدم من الفلتر
        states = [state for state in (states or []) if state in ('draft', 'incomplete', 'posted')] or ['posted']
        moves = self.search([
            ('ledger_year', '=', int(year)),
            ('ledger_month', '=', str(int(month))),
            ('state', 'in', states),
        ], order='sort_key, id')

        # ترتيب الأعمدة حسب "ترتيب في دفتر الأستاذ" المحدَّد في كل حساب رئيسي
        accounts = self.env['myaccounting.account'].search(
            [('parent_id', '=', False)], order='ledger_sequence, code')

        rows = []
        totals = {acc.id: {'debit': 0.0, 'credit': 0.0, 'debit_zero': False, 'credit_zero': False}
                  for acc in accounts}
        grand_debit = 0.0
        grand_credit = 0.0

        for idx, move in enumerate(moves, start=1):
            amounts = {}
            for line in move.line_ids:
                if not line.account_id:
                    # بند في قيد "غير مكتمل" لم يُحدَّد حسابه بعد: لا عمود له،
                    # لكن مبلغه يظهر ضمن إجمالي مدين/دائن للقيد.
                    continue
                root_id = int((line.account_id.parent_path or str(line.account_id.id)).split('/')[0])
                if root_id not in amounts:
                    amounts[root_id] = {'debit': 0.0, 'credit': 0.0, 'debit_zero': False, 'credit_zero': False}
                amounts[root_id]['debit'] += line.debit
                amounts[root_id]['credit'] += line.credit
                # صفر مُدخل يدوياً (مثل ضريبة "معفي"): يُعرض 0 بدل خانة فارغة
                amounts[root_id]['debit_zero'] |= line.debit_zero_entered
                amounts[root_id]['credit_zero'] |= line.credit_zero_entered
                if root_id in totals:
                    totals[root_id]['debit'] += line.debit
                    totals[root_id]['credit'] += line.credit
                    totals[root_id]['debit_zero'] |= line.debit_zero_entered
                    totals[root_id]['credit_zero'] |= line.credit_zero_entered
            rows.append({
                'seq': idx,
                'move_id': move.id,
                'move_name': move.name,
                'state': move.state,
                'ref': move.ref or '',
                'date': move.date and move.date.isoformat(),
                'total_debit': move.total_debit,
                'total_credit': move.total_credit,
                'amounts': amounts,
            })
            grand_debit += move.total_debit
            grand_credit += move.total_credit

        return {
            'accounts': [{'id': a.id, 'code': a.code, 'name': a.name} for a in accounts],
            'rows': rows,
            'totals': {a.id: totals[a.id] for a in accounts},
            'grand_debit': grand_debit,
            'grand_credit': grand_credit,
        }

    # ========================================================================
    # الاستيراد من ملف Excel
    # ========================================================================

    # ترويسات الأعمدة المتوقّعة في الملف، وما يقابلها داخلياً
    IMPORT_COLUMNS = {
        'رقم القيد': 'move_key',
        'التاريخ': 'date',
        'المرجع': 'ref',
        'اليومية': 'journal',
        'شهر دفتر الأستاذ': 'ledger_month',
        'سنة دفتر الأستاذ': 'ledger_year',
        'رمز الحساب': 'account_code',
        'اسم الحساب': 'account_name',
        'البيان': 'label',
        'مدين': 'debit',
        'دائن': 'credit',
    }

    @api.model
    def build_import_template_xlsx(self):
        """قالب Excel فارغ لاستيراد القيود (مع سطري مثال)."""
        return self._build_moves_xlsx()

    @api.model
    def export_moves_to_xlsx(self, move_ids):
        """يصدّر القيود المحددة إلى ملف Excel بنفس قالب الاستيراد وبنفس ترتيب
        القائمة المعروضة، فيمكن تعديله أو الإضافة عليه وإعادة رفعه."""
        moves = self.browse(move_ids).exists()
        position = {move_id: index for index, move_id in enumerate(move_ids)}
        moves = moves.sorted(key=lambda move: position.get(move.id, 0))
        return base64.b64encode(self._build_moves_xlsx(moves)).decode()

    @api.model
    def _build_moves_xlsx(self, moves=None):
        """يبني ملف Excel بنفس القالب: ورقة القيود (مثال أو قيود فعلية)،
        ورقة تعليمات، وورقة بكل الحسابات المتاحة."""
        import xlsxwriter

        output = io.BytesIO()
        workbook = xlsxwriter.Workbook(output, {'in_memory': True})

        header_fmt = workbook.add_format({
            'bold': True, 'bg_color': '#D9E1F2', 'border': 1, 'text_wrap': True, 'align': 'center',
        })
        text_fmt = workbook.add_format({'border': 1})
        num_fmt = workbook.add_format({'border': 1, 'num_format': '#,##0.000'})
        date_fmt = workbook.add_format({'border': 1, 'num_format': 'yyyy-mm-dd'})
        note_fmt = workbook.add_format({'text_wrap': True, 'valign': 'top'})
        title_fmt = workbook.add_format({'bold': True, 'font_size': 13})

        # --- ورقة القيود ---
        sheet = workbook.add_worksheet('القيود')
        sheet.right_to_left()
        headers = [
            'رقم القيد', 'التاريخ', 'المرجع', 'اليومية',
            'شهر دفتر الأستاذ', 'سنة دفتر الأستاذ',
            'رمز الحساب', 'اسم الحساب', 'البيان', 'مدين', 'دائن',
        ]
        for col, header in enumerate(headers):
            sheet.write(0, col, header, header_fmt)

        today = fields.Date.context_today(self)
        sample_accounts = self.env['myaccounting.account'].search(
            [('parent_id', '=', False)], order='code', limit=2)
        first_name = sample_accounts[0].name if sample_accounts else 'اسم حساب من شجرة الحسابات'
        second_name = sample_accounts[1].name if len(sample_accounts) > 1 else 'اسم حساب آخر'

        if moves is None:
            data_rows = [
                ['1', today, 'مرجع اختياري', 'القيود اليدوية', today.month, today.year,
                 '', first_name, 'شرح الحركة', 500, ''],
                ['', '', '', '', '', '', '', second_name, 'شرح الحركة', '', 500],
            ]
        else:
            # كل قيد: بيانات القيد في سطر بنده الأول، ثم بقية البنود بأعمدة قيد فارغة
            data_rows = []
            for move in moves:
                header = [move.name or '', move.date or '', move.ref or '',
                          move.journal or '', int(move.ledger_month or 0) or '', move.ledger_year or '']
                lines = move.line_ids or [None]
                for index, line in enumerate(lines):
                    prefix = header if index == 0 else ['', '', '', '', '', '']
                    if line is None:
                        data_rows.append(prefix + ['', '', '', '', ''])
                        continue
                    data_rows.append(prefix + [
                        line.account_id.code or '',
                        line.account_id.name or line.pending_account_name or '',
                        line.name or '',
                        line.debit or (0 if line.debit_zero_entered else ''),
                        line.credit or (0 if line.credit_zero_entered else ''),
                    ])
        for row_index, row in enumerate(data_rows, start=1):
            for col, value in enumerate(row):
                if col == 1 and value:
                    sheet.write_datetime(row_index, col, value, date_fmt)
                elif col in (9, 10) and value != '':
                    sheet.write_number(row_index, col, value, num_fmt)
                else:
                    sheet.write(row_index, col, value, text_fmt)

        sheet.set_column(0, 0, 10)
        sheet.set_column(1, 1, 12)
        sheet.set_column(2, 3, 16)
        sheet.set_column(4, 5, 16)
        sheet.set_column(6, 6, 12)
        sheet.set_column(7, 8, 24)
        sheet.set_column(9, 10, 13)
        sheet.freeze_panes(1, 0)

        # --- ورقة التعليمات ---
        guide = workbook.add_worksheet('تعليمات')
        guide.right_to_left()
        guide.set_column(0, 0, 100)
        guide.write(0, 0, 'طريقة تعبئة الملف', title_fmt)
        instructions = [
            '1) كل سطر يمثّل بنداً واحداً داخل القيد (حساب واحد بمبلغ مدين أو دائن).',
            '2) البنود التي تحمل نفس "رقم القيد" تُجمَّع في قيد واحد.',
            '   يكفي كتابة رقم القيد في أول سطر، والأسطر التالية التي يكون فيها العمود فارغاً تتبع نفس القيد.',
            '3) يمكن كتابة مبلغ في "مدين" و"دائن" معاً في نفس السطر عند الحاجة، ولا تُقبل المبالغ السالبة.',
            '4) "التاريخ" و"المرجع" و"اليومية" و"شهر/سنة دفتر الأستاذ" تُؤخذ من أول سطر في كل قيد.',
            '   إذا تُرك التاريخ أو شهر/سنة دفتر الأستاذ فارغاً يُؤخذ من القيد السابق في الملف',
            '   (ولأول قيد: تاريخ اليوم، وشهر/سنة دفتر الأستاذ من التاريخ).',
            '   التاريخ يُكتب اليوم أولاً بأي شكل: 25/7/2026 أو 25-7-26 أو 25.07.2026 أو 2026-07-25،',
            '   ويمكن كتابة اليوم والشهر فقط (25/7) فتُؤخذ السنة من "سنة دفتر الأستاذ" أو من القيد السابق.',
            '   شهر دفتر الأستاذ يُكتب رقماً (7) أو اسماً (يوليو / تموز).',
            '5) "رمز الحساب" اختياري — إذا كتبته تتم المطابقة به أولاً، وإلّا فبـ"اسم الحساب".',
            '6) إذا لم يُعثر على الحساب في شجرة الحسابات (أو كان الاسم مكرّراً في أكثر من حساب)،',
            '   يُستورد القيد بحالة "غير مكتمل" مع حفظ اسم الحساب كما ورد في الملف.',
            '7) بعد الاستيراد: افتح صفحة القيود، اضغط فلتر الحالة "غير مكتمل"، ثم افتح القيد واختر الحساب الصحيح.',
            '   بمجرد اختيارك الحساب لبند واحد، تُحدَّث تلقائياً كل البنود غير المكتملة التي تحمل نفس اسم الحساب.',
            '8) رقم القيد: إذا كان مستخدماً مسبقاً في النظام، يُعطى القيد رقماً جديداً تلقائياً.',
        ]
        if moves is None:
            instructions.append('9) احذف سطري المثال قبل رفع الملف.')
        else:
            instructions.append('9) هذا الملف مُصدَّر من النظام: عدّل عليه أو أضف قيوداً جديدة ثم ارفعه من زر "استيراد من Excel".')
            instructions.append('   انتبه: إعادة رفع قيد رقمه موجود في النظام تُنشئ قيداً جديداً برقم جديد ولا تُعدّل القيد الأصلي.')
        for index, line in enumerate(instructions, start=2):
            guide.write(index, 0, line, note_fmt)

        # --- ورقة الحسابات المتاحة ---
        accounts_sheet = workbook.add_worksheet('الحسابات المتاحة')
        accounts_sheet.right_to_left()
        for col, header in enumerate(['رمز الحساب', 'اسم الحساب', 'الحساب الأب']):
            accounts_sheet.write(0, col, header, header_fmt)
        for index, account in enumerate(
                self.env['myaccounting.account'].search([], order='code_path'), start=1):
            accounts_sheet.write(index, 0, account.code or '', text_fmt)
            accounts_sheet.write(index, 1, account.name or '', text_fmt)
            accounts_sheet.write(index, 2, account.parent_id.name or '', text_fmt)
        accounts_sheet.set_column(0, 0, 12)
        accounts_sheet.set_column(1, 2, 30)
        accounts_sheet.freeze_panes(1, 0)

        workbook.close()
        output.seek(0)
        return output.read()

    @api.model
    def _import_parse_date(self, value, default_year=None):
        """يحوّل قيمة خلية إلى تاريخ. يُرجع None إن كانت فارغة، و False إن تعذّر التحويل.

        الأشكال المقبولة (اليوم أولاً دائماً، وأي فاصل من / - . أو مسافة):
          25/7/2026 ، 25-7-26 ، 25.07.2026 ، 2026-07-25 ، ٢٥/٧/٢٠٢٦
          25/7 (بدون سنة) ← تُؤخذ السنة من default_year
          رقم تاريخ Excel التسلسلي (مثل 46228)
        """
        if value in (None, ''):
            return None
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            # خلية رقمية: رقم تاريخ Excel التسلسلي (أيام منذ 1899-12-30)
            if 20000 <= value <= 80000:
                return date(1899, 12, 30) + timedelta(days=int(value))
            return False
        text = normalize_digits(str(value)).strip()
        # "2026-07-25 00:00:00" ← نتجاهل الوقت إن وُجد
        text = re.sub(r'\s+\d{1,2}:\d{2}(:\d{2})?$', '', text)
        if not text:
            return None
        if text.isdigit() and 20000 <= int(text) <= 80000:
            return date(1899, 12, 30) + timedelta(days=int(text))
        parts = re.split(r'\s*[/\-.\\\s]\s*', text)
        if not all(part.isdigit() for part in parts):
            return False
        numbers = [int(part) for part in parts]
        if len(parts) == 3:
            if len(parts[0]) == 4:
                year, month, day = numbers
            else:
                day, month, year = numbers
                if len(parts[2]) <= 2:
                    year += 2000
        elif len(parts) == 2 and len(parts[0]) <= 2 and len(parts[1]) <= 2:
            day, month = numbers
            year = default_year or fields.Date.context_today(self).year
        else:
            return False
        try:
            return date(year, month, day)
        except ValueError:
            return False

    @api.model
    def _import_parse_ledger_value(self, value, kind):
        """يحوّل خلية شهر/سنة دفتر الأستاذ إلى رقم صحيح.
        يُرجع None إن كانت فارغة، و False إن تعذّر التحويل.
        يقبل: 7 ، 07 ، 7.0 ، ٧ ، أسماء الأشهر (يوليو / تموز)، والسنة بخانتين (26 ← 2026)."""
        if value in (None, ''):
            return None
        if isinstance(value, bool):
            return False
        if isinstance(value, float):
            if not value.is_integer():
                return False
            value = int(value)
        if isinstance(value, int):
            number = value
        else:
            text = normalize_digits(str(value)).strip()
            if not text:
                return None
            if kind == 'month':
                month = ARABIC_MONTHS.get(normalize_account_key(text))
                if month:
                    return month
            if re.fullmatch(r'\d+(\.0+)?', text):
                number = int(float(text))
            else:
                return False
        if kind == 'year':
            if 0 <= number < 100:
                number += 2000
            return number if 1900 <= number <= 2200 else False
        return number if 1 <= number <= 12 else False

    @api.model
    def _import_parse_float(self, value):
        """يحوّل قيمة خلية إلى رقم. يُرجع None إن تعذّر التحويل."""
        if value in (None, ''):
            return 0.0
        if isinstance(value, bool):
            return 0.0
        if isinstance(value, (int, float)):
            return float(value)
        text = str(value).strip().replace(',', '')
        if not text:
            return 0.0
        try:
            return float(text)
        except ValueError:
            return None

    @api.model
    def _import_build_account_index(self):
        """يبني فهرسين للحسابات: حسب الرمز وحسب الاسم.
        الأسماء المتكرّرة تُعلَّم كـ None ليُطلب اختيارها يدوياً لاحقاً."""
        accounts = self.env['myaccounting.account'].search([])
        by_code, by_name = {}, {}
        for account in accounts:
            code_key = normalize_account_key(account.code)
            if code_key:
                by_code[code_key] = account.id
            name_key = normalize_account_key(account.name)
            if name_key:
                by_name[name_key] = None if name_key in by_name else account.id
        return by_code, by_name

    @api.model
    def _import_next_names(self, count, used):
        """يولّد أرقام قيود متسلسلة غير مستخدمة (داخل قاعدة البيانات أو ضمن نفس الاستيراد)."""
        names = []
        last_move = self.search([], order='id desc', limit=1)
        candidate = last_move.name if last_move and last_move.name else '0'
        for _ in range(count):
            guard = 0
            while guard < 10000:
                candidate = self._increment_name(candidate)
                if candidate not in used and not self.search_count([('name', '=', candidate)]):
                    break
                guard += 1
            used.add(candidate)
            names.append(candidate)
        return names

    def _post_import_notes(self, notes):
        """يحفظ ملاحظات الاستيراد على القيد وينشرها كـ"ملاحظة" داخلية في المحادثة."""
        self.ensure_one()
        items = Markup('').join(Markup('<li>%s</li>') % escape(note) for note in notes)
        body = Markup('<p><b>ملاحظات الاستيراد من Excel (%s)</b></p><ul>%s</ul>') % (len(notes), items)
        self.write({'import_notes': body, 'import_notes_reviewed': False})
        self.message_post(body=body, subtype_xmlid='mail.mt_note')

    @api.model
    def import_moves_from_xlsx(self, file_b64):
        """يستورد قيوداً متعدّدة من ملف Excel.

        تُجمَّع الأسطر التي تحمل نفس "رقم القيد" في قيد واحد. أي بند لم يُعثر على
        حسابه (أو كان اسمه مكرّراً في شجرة الحسابات) يُحفظ باسم الحساب كما ورد في
        الملف، ويأخذ القيد حالة "غير مكتمل" ليُستكمل يدوياً لاحقاً.
        """
        try:
            import openpyxl
        except ImportError:
            raise UserError('مكتبة قراءة ملفات Excel غير متوفرة على الخادم (openpyxl).')

        try:
            data = base64.b64decode(file_b64)
            workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        except Exception:
            raise UserError('تعذّرت قراءة الملف. تأكد أنه ملف Excel بصيغة .xlsx سليم.')

        sheet = workbook.worksheets[0]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            raise UserError('الملف فارغ.')

        # ربط ترويسات الأعمدة بمواقعها
        header_map = {}
        for index, cell in enumerate(rows[0]):
            key = self.IMPORT_COLUMNS.get(' '.join(str(cell or '').split()))
            if key:
                header_map[key] = index
        for required in ('move_key', 'account_name', 'debit', 'credit'):
            if required not in header_map:
                raise UserError(
                    'ترويسات الأعمدة غير مطابقة للقالب. الأعمدة المطلوبة على الأقل: '
                    'رقم القيد، اسم الحساب، مدين، دائن.\n'
                    'حمّل القالب من الزر المخصّص واستخدمه كما هو.'
                )

        def cell(row, key):
            index = header_map.get(key)
            if index is None or index >= len(row):
                return None
            return row[index]

        by_code, by_name = self._import_build_account_index()
        Account = self.env['myaccounting.account']

        # تجميع الأسطر في قيود: السطر بلا "رقم قيد" يتبع القيد السابق
        groups, order, current_key = {}, [], None
        for row_number, row in enumerate(rows[1:], start=2):
            if row is None or all(value in (None, '') for value in row):
                continue
            raw_key = ' '.join(str(cell(row, 'move_key') or '').split())
            if raw_key:
                current_key = raw_key
            if current_key is None:
                current_key = f'__بلا_رقم_{row_number}'
            if current_key not in groups:
                groups[current_key] = []
                order.append(current_key)
            groups[current_key].append((row_number, row))

        created, incomplete_count, errors = [], 0, []
        unmatched_names = {}
        used_names = set()
        generated_names = iter(self._import_next_names(len(order), used_names))
        previous = None  # قيم القيد السابق في الملف (التاريخ وشهر/سنة دفتر الأستاذ)

        for key in order:
            group = groups[key]
            first_row = group[0][1]
            entry_errors = []

            # entry_errors: كل خطأ أو تصحيح تلقائي يخص هذا القيد، يُحفظ معه ويُنشر كملاحظة.
            # الخلايا الفارغة ليست خطأ: التاريخ وشهر/سنة دفتر الأستاذ يُؤخذان من القيد
            # السابق في الملف (فالغالب أنها نفس القيم)، دون تسجيل ملاحظة.
            today = fields.Date.context_today(self)
            raw_date = cell(first_row, 'date')
            raw_year = cell(first_row, 'ledger_year')
            raw_month = cell(first_row, 'ledger_month')

            ledger_year = self._import_parse_ledger_value(raw_year, 'year')
            ledger_month = self._import_parse_ledger_value(raw_month, 'month')

            # سنة التاريخ المختصر (25/7): سنة دفتر الأستاذ، ثم سنة القيد السابق، ثم السنة الحالية
            default_year = ledger_year or (previous['date'].year if previous else today.year)
            move_date = self._import_parse_date(raw_date, default_year=default_year)
            if not move_date:
                fallback = previous['date'] if previous else today
                fallback_label = 'تاريخ القيد السابق' if previous else 'تاريخ اليوم'
                if move_date is False:
                    entry_errors.append(f'تاريخ غير صالح: "{raw_date}" ← استُخدم {fallback_label} {fallback}')
                elif not previous:
                    entry_errors.append(f'لا يوجد تاريخ في الملف ← استُخدم تاريخ اليوم {today}')
                move_date = fallback

            if not ledger_year:
                fallback = previous['ledger_year'] if previous else move_date.year
                if ledger_year is False:
                    source = 'سنة القيد السابق' if previous else 'سنة التاريخ'
                    entry_errors.append(f'سنة دفتر أستاذ غير صالحة: "{raw_year}" ← استُخدمت {source} {fallback}')
                ledger_year = fallback

            if not ledger_month:
                fallback = previous['ledger_month'] if previous else move_date.month
                if ledger_month is False:
                    source = 'شهر القيد السابق' if previous else 'شهر التاريخ'
                    entry_errors.append(f'شهر دفتر أستاذ غير صالح: "{raw_month}" ← استُخدم {source} {fallback}')
                ledger_month = fallback

            previous = {'date': move_date, 'ledger_year': ledger_year, 'ledger_month': ledger_month}

            line_commands, has_pending = [], False
            for row_number, row in group:
                account_name = ' '.join(str(cell(row, 'account_name') or '').split())
                account_code = ' '.join(str(cell(row, 'account_code') or '').split())
                debit = self._import_parse_float(cell(row, 'debit'))
                credit = self._import_parse_float(cell(row, 'credit'))
                # 0 مكتوب في الخلية (وليس خلية فارغة)
                debit_zero = debit == 0 and cell(row, 'debit') not in (None, '') \
                    and str(cell(row, 'debit')).strip() != ''
                credit_zero = credit == 0 and cell(row, 'credit') not in (None, '') \
                    and str(cell(row, 'credit')).strip() != ''

                if debit is None or credit is None:
                    entry_errors.append(f'سطر {row_number}: قيمة مدين/دائن غير رقمية ← لم يُستورد السطر')
                    continue
                if not account_name and not account_code and not debit and not credit:
                    continue
                if not account_name and not account_code:
                    entry_errors.append(f'سطر {row_number}: لا يوجد اسم أو رمز حساب ← لم يُستورد السطر')
                    continue
                if debit < 0 or credit < 0:
                    entry_errors.append(f'سطر {row_number}: لا يمكن أن يكون المبلغ سالباً ← لم يُستورد السطر')
                    continue

                account_id = by_code.get(normalize_account_key(account_code)) if account_code else None
                if account_id:
                    account = Account.browse(account_id)
                    if account_name and normalize_account_key(account_name) != normalize_account_key(account.name):
                        entry_errors.append(
                            f'سطر {row_number}: الرمز "{account_code}" يعود للحساب "{account.name}" '
                            f'بينما الاسم في الملف "{account_name}" ← اعتُمد الحساب حسب الرمز')
                elif account_name:
                    name_key = normalize_account_key(account_name)
                    account_id = by_name.get(name_key)
                    if account_id:
                        if account_code:
                            account = Account.browse(account_id)
                            entry_errors.append(
                                f'سطر {row_number}: الرمز "{account_code}" غير موجود ← تمت المطابقة بالاسم '
                                f'مع الحساب {account.code} - {account.name}')
                    elif name_key in by_name:
                        entry_errors.append(
                            f'سطر {row_number}: اسم الحساب "{account_name}" مكرّر في شجرة الحسابات '
                            f'← يجب اختيار الحساب يدوياً')
                    else:
                        entry_errors.append(
                            f'سطر {row_number}: الحساب "{account_name}"'
                            + (f' (رمز "{account_code}")' if account_code else '')
                            + ' غير موجود ← يجب اختياره أو إنشاؤه يدوياً')
                else:
                    entry_errors.append(
                        f'سطر {row_number}: الرمز "{account_code}" غير موجود ولا يوجد اسم حساب '
                        f'← يجب اختيار الحساب يدوياً')

                values = {
                    'name': (str(cell(row, 'label')).strip() if cell(row, 'label') not in (None, '') else False),
                    'debit': debit,
                    'credit': credit,
                    'debit_zero_entered': debit_zero,
                    'credit_zero_entered': credit_zero,
                }
                if account_id:
                    values['account_id'] = account_id
                else:
                    label = account_name or account_code
                    values['pending_account_name'] = label
                    has_pending = True
                    unmatched_names[label] = unmatched_names.get(label, 0) + 1
                line_commands.append((0, 0, values))

            if not line_commands:
                if entry_errors:
                    errors.append(f'القيد "{key}": ' + '؛ '.join(entry_errors))
                continue

            file_name = key if key and not key.startswith('__بلا_رقم_') else None
            move_name = file_name
            if move_name and (move_name in used_names or self.search_count([('name', '=', move_name)])):
                move_name = None
            if move_name:
                used_names.add(move_name)
            else:
                move_name = next(generated_names)
                if file_name:
                    entry_errors.insert(0, f'رقم القيد "{file_name}" مستخدم مسبقاً ← أُعطي الرقم {move_name}')
                else:
                    entry_errors.insert(0, f'لا يوجد رقم قيد في الملف ← أُعطي الرقم {move_name}')

            move = self.create({
                'name': move_name,
                'date': move_date,
                'ref': (str(cell(first_row, 'ref')).strip() if cell(first_row, 'ref') not in (None, '') else False),
                'journal': (str(cell(first_row, 'journal')).strip()
                            if cell(first_row, 'journal') not in (None, '') else 'القيود اليدوية'),
                'ledger_month': str(ledger_month),
                'ledger_year': ledger_year,
                'state': 'incomplete' if has_pending else 'draft',
                'line_ids': line_commands,
            })
            created.append(move)
            if has_pending:
                incomplete_count += 1
            if not move.is_balanced:
                entry_errors.append(
                    f'القيد غير متوازن: مدين {move.total_debit:,.3f} / دائن {move.total_credit:,.3f} '
                    f'(الفرق {abs(move.total_debit - move.total_credit):,.3f}) ← يجب تصحيحه قبل الترحيل')
            if entry_errors:
                errors.append(f'القيد "{move.name}": ' + '؛ '.join(entry_errors))
                move._post_import_notes(entry_errors)

        unbalanced = [move.name for move in created if not move.is_balanced]

        return {
            'created': len(created),
            'incomplete': incomplete_count,
            'unbalanced': unbalanced,
            'unmatched_accounts': sorted(unmatched_names.items(), key=lambda item: -item[1]),
            'errors': errors,
        }


class MyAccountingMoveLine(models.Model):
    _name = 'myaccounting.move.line'
    _description = 'بند قيد محاسبي'
    # الترتيب اليدوي للبنود (سحب وإفلات) ثم ترتيب الإدخال للبنود المتساوية
    _order = 'sequence, id'

    sequence = fields.Integer(string='الترتيب', default=10)

    move_id = fields.Many2one('myaccounting.move', string='القيد', required=True, ondelete='cascade')
    move_state = fields.Selection(related='move_id.state', string='حالة القيد', store=True)
    move_sort_key = fields.Char(related='move_id.sort_key', string='ترتيب القيد', store=True, index=True)
    account_id = fields.Many2one('myaccounting.account', string='الحساب')
    pending_account_name = fields.Char(
        string='اسم الحساب في الملف', copy=False,
        help='اسم الحساب كما ورد في ملف Excel عندما تعذّر مطابقته مع شجرة الحسابات. '
             'عند اختيار الحساب الصحيح هنا، تُحدَّث تلقائياً كل البنود غير المكتملة التي تحمل نفس الاسم.',
    )
    account_label = fields.Char(string='اسم الحساب المعروض في القيد', compute='_compute_account_label')
    name = fields.Char(string='البيان')
    debit = fields.Float(string='مدين', default=0.0, digits=(16, 3))
    credit = fields.Float(string='دائن', default=0.0, digits=(16, 3))
    # الخانة الفارغة تُخزَّن 0 أيضاً، لكن الصفر المُدخل يدوياً له معنى محاسبي آخر:
    # هذان العلمان يميّزانه، فيُعرض "0" فقط عندما أُدخل فعلاً ويبقى غيره فارغاً.
    debit_zero_entered = fields.Boolean(string='صفر مُدخل في المدين')
    credit_zero_entered = fields.Boolean(string='صفر مُدخل في الدائن')
    currency_id = fields.Many2one(related='move_id.currency_id', string='العملة', store=True)
    date = fields.Date(related='move_id.date', string='التاريخ', store=True)

    # عرض القيد بعمودين للحساب: يظهر الحساب تحت "الحساب المدين" أو "الحساب الدائن"
    # حسب القيمة المدخلة (قبل إدخال القيمة يظهر في المدين). الحقل الفعلي يبقى account_id.
    debit_account_id = fields.Many2one(
        'myaccounting.account', string='الحساب المدين',
        compute='_compute_side_accounts', inverse='_inverse_debit_account')
    credit_account_id = fields.Many2one(
        'myaccounting.account', string='الحساب الدائن',
        compute='_compute_side_accounts', inverse='_inverse_credit_account')
    is_credit_line = fields.Boolean(string='بند دائن', compute='_compute_side_accounts')

    @api.depends('account_id', 'debit', 'credit')
    def _compute_side_accounts(self):
        for line in self:
            is_credit = bool(line.credit and not line.debit)
            line.is_credit_line = is_credit
            line.debit_account_id = False if is_credit else line.account_id
            line.credit_account_id = line.account_id if is_credit else False

    # القيمة الفارغة لا تُكتب: عند انتقال الحساب من عمود لآخر يصبح أحدهما فارغاً
    # بالحساب لا بقرار المستخدم، فلا يجب أن يمسح الحساب الفعلي.
    def _inverse_debit_account(self):
        for line in self:
            if line.debit_account_id:
                line.account_id = line.debit_account_id

    def _inverse_credit_account(self):
        for line in self:
            if line.credit_account_id:
                line.account_id = line.credit_account_id

    @api.onchange('debit_account_id')
    def _onchange_debit_account_id(self):
        if self.debit_account_id:
            self.account_id = self.debit_account_id

    @api.onchange('credit_account_id')
    def _onchange_credit_account_id(self):
        if self.credit_account_id:
            self.account_id = self.credit_account_id

    @api.depends('account_id', 'account_id.name', 'account_id.parent_id.name', 'pending_account_name')
    def _compute_account_label(self):
        # عند تكرار اسم الحساب في أكثر من مكان في شجرة الحسابات (مثال: "الرواتب"
        # تحت أكثر من حساب أب)، نضيف اسم الحساب الأب لتمييزه، لكن فقط في عرض
        # بند القيد هذا، دون المساس باسم الحساب نفسه أو عرضه في أي مكان آخر.
        for line in self:
            account = line.account_id
            if not account:
                line.account_label = line.pending_account_name or ''
                continue
            is_duplicate = bool(self.env['myaccounting.account'].search_count([
                ('name', '=', account.name),
                ('id', '!=', account.id),
            ]))
            if is_duplicate and account.parent_id:
                line.account_label = f"{account.name} / {account.parent_id.name}"
            else:
                line.account_label = account.name

    @api.model_create_multi
    def create(self, vals_list):
        lines = super().create(vals_list)
        lines.move_id._sync_state_from_lines()
        return lines

    def write(self, vals):
        res = super().write(vals)
        if 'account_id' in vals and vals.get('account_id'):
            self._propagate_account_to_matching_lines()
        if 'account_id' in vals:
            self.move_id._sync_state_from_lines()
        return res

    def unlink(self):
        moves = self.move_id
        res = super().unlink()
        moves.exists()._sync_state_from_lines()
        return res

    def _propagate_account_to_matching_lines(self):
        """عند اختيار حساب لبند كان اسمه غير معروف، يُطبَّق نفس الحساب تلقائياً على
        كل البنود الأخرى غير المكتملة التي تحمل نفس اسم الحساب في الملف."""
        for line in self:
            pending = (line.pending_account_name or '').strip()
            if not pending or not line.account_id:
                continue

            key = normalize_account_key(pending)
            candidates = self.search([
                ('id', '!=', line.id),
                ('account_id', '=', False),
                ('pending_account_name', '!=', False),
            ])
            matching = candidates.filtered(
                lambda other: normalize_account_key(other.pending_account_name) == key
            )

            line.pending_account_name = False
            if matching:
                matching.write({
                    'account_id': line.account_id.id,
                    'pending_account_name': False,
                })
                matching.move_id._sync_state_from_lines()

    @api.constrains('debit', 'credit')
    def _check_debit_credit(self):
        for line in self:
            if line.debit < 0 or line.credit < 0:
                raise ValidationError('لا يمكن أن تكون قيمة المدين أو الدائن سالبة.')

    def action_reorder(self, delta):
        """يحرّك البند خطوة واحدة (delta = -1 لأعلى، 1 لأسفل) داخل قيده.

        يعمل حتى على القيد المرحّل لأن الترتيب عرضي فقط ولا يغيّر المبالغ."""
        self.ensure_one()
        lines = list(self.move_id.line_ids)
        index = lines.index(self)
        target = index + delta
        if target < 0 or target >= len(lines):
            return False
        lines.insert(target, lines.pop(index))
        for position, line in enumerate(lines, start=1):
            if line.sequence != position:
                line.sequence = position
        return True
